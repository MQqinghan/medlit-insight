"""报告导出：Markdown（文本）与 Word（python-docx）。PDF 走浏览器打印，见前端 printReport()。"""
from datetime import datetime
from io import BytesIO

DOCX_AVAILABLE = True
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover
    DOCX_AVAILABLE = False

DISCLAIMER = (
    "数据来源：PubMed E-utilities（检索时间 {ts}）。期刊影响因子与分区为本地静态示例库数据"
    "（版本 {ver}），非实时权威口径；{llm}"
)


def _llm_note(result):
    """按综述实际来源生成声明：LLM 生成时标注模型名，兜底时明示模板生成。"""
    review = result.get("review", {})
    if review.get("source") == "llm":
        return f"综述由 {review.get('model') or 'LLM'} 基于检索摘要生成，仅供参考。"
    return "综述为统计模板自动生成（本次未调用 LLM），仅供参考。"


def _fmt_pct(x):
    return "—" if x is None else f"{round(x * 100)}%"


def build_markdown(result):
    s = result["stats"]
    p = result["params"]
    ts = result.get("generated_at", "")
    ver = result.get("meta", {}).get("journal_db_version", "—")
    lines = [
        f"# 文献洞察报告：{result['keyword']}",
        "",
        f"- 生成时间：{ts}",
        f"- 检索范围：近{p['years']}年（{p['start_year']}–{p['end_year']}），分析样本上限 {p['max_fetch']} 篇",
        f"- 检索命中总量：{s['total_count']} 篇；近{p['years']}年：{s['recent_count']} 篇（{_fmt_pct(s['recent_ratio'])}）",
        f"- 纳入分析：{s['analyzed_count']} 篇；期刊指标匹配率：{_fmt_pct(s['matched_ratio'])}",
        f"- 样本中位 IF：{s['median_if'] if s['median_if'] is not None else '—'}；1区期刊占比：{_fmt_pct(s['q1_ratio'])}",
        "",
        "## 一、领域综述",
        "",
        result["review"]["markdown"],
        "",
        "## 二、研究热点词云（MeSH 词频 Top30）",
        "",
        "、".join(w["name"] for w in result["wordcloud"][:30]) or "—",
        "",
        "## 三、近5年影响力 Top{} 文献".format(len(s["top_n"])),
        "",
        "| 排名 | 标题 | 期刊 | 年份 | IF | 分区 |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for a in s["top_n"]:
        title = a["title"].replace("|", "\\|")
        lines.append(
            f"| {a['rank']} | {title} | {a['journal']} | {a['year']} | {a['if']} | {a['zone'] or '—'} |"
        )
    lines += ["", DISCLAIMER.format(ts=ts, ver=ver, llm=_llm_note(result)), ""]
    return "\n".join(lines)


def build_docx(result):
    if not DOCX_AVAILABLE:
        raise RuntimeError("python-docx 未安装")
    s = result["stats"]
    p = result["params"]
    ver = result.get("meta", {}).get("journal_db_version", "—")
    ts = result.get("generated_at", "")

    doc = Document()
    title = doc.add_heading(f"文献洞察报告：{result['keyword']}", level=0)

    meta = doc.add_paragraph()
    meta.add_run(
        f"生成时间：{ts}\n检索范围：近{p['years']}年（{p['start_year']}–{p['end_year']}）"
        f"\n命中总量：{s['total_count']} 篇；近{p['years']}年：{s['recent_count']} 篇"
        f"\n纳入分析：{s['analyzed_count']} 篇；中位 IF：{s['median_if'] if s['median_if'] is not None else '—'}"
        f"\n1区期刊占比：{_fmt_pct(s['q1_ratio'])}；期刊指标匹配率：{_fmt_pct(s['matched_ratio'])}"
    )
    meta.runs[0].font.size = Pt(10)

    doc.add_heading("一、领域综述", level=1)
    for para in result["review"]["markdown"].split("\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_heading("二、研究热点词云（MeSH 词频 Top30）", level=1)
    doc.add_paragraph("、".join(w["name"] for w in result["wordcloud"][:30]) or "—")

    doc.add_heading(f"三、近5年影响力 Top{len(s['top_n'])} 文献", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["排名", "标题", "期刊", "年份", "IF", "分区"]):
        table.rows[0].cells[i].text = h
    for a in s["top_n"]:
        row = table.add_row().cells
        row[0].text = str(a["rank"])
        row[1].text = a["title"]
        row[2].text = a["journal"]
        row[3].text = str(a["year"])
        row[4].text = str(a["if"])
        row[5].text = a["zone"] or "—"

    note = doc.add_paragraph()
    run = note.add_run(DISCLAIMER.format(ts=ts, ver=ver, llm=_llm_note(result)))
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
